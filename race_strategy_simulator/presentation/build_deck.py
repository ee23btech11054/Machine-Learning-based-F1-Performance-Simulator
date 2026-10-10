"""
Build the project presentation (.pptx), mathematical version.

Numbers come from the project's result files and from
presentation/assets/numbers.json (written by deck_analysis.py). Equations are
typeset with matplotlib and placed as images.

    python presentation/deck_analysis.py
    python presentation/build_deck.py <project folder> <output .pptx>
"""
import hashlib
import json
import math
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

project = Path(sys.argv[1])
output = Path(sys.argv[2])
reports = project / "reports"
assets = project / "presentation" / "assets"
equations = assets / "equations"
equations.mkdir(parents=True, exist_ok=True)

validation = pd.read_csv(reports / "03_models" / "validation_comparison.csv")
test = pd.read_csv(reports / "03_models" / "test_results.csv").set_index("model")
backtest = pd.read_csv(reports / "04_backtest" / "backtest_summary.csv").set_index("metric")["value"]
num = json.loads((assets / "numbers.json").read_text())

MODEL = "Gradient boosting"
NAIVE = "No change from baseline"

# ------------------------------------------------------------------ style ---
INK = RGBColor(0x14, 0x17, 0x1A)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xE1, 0x06, 0x00)
TEXT = RGBColor(0x1E, 0x22, 0x26)
MUTED = RGBColor(0x5B, 0x61, 0x68)
PANEL = RGBColor(0xF2, 0xF3, 0xF5)
GREY = RGBColor(0xB4, 0xB9, 0xC0)
MID = RGBColor(0x6E, 0x75, 0x7D)
SOFT_WHITE = RGBColor(0xC9, 0xCD, 0xD2)

HEAD = "Arial"
BODY = "Calibri"

deck = Presentation()
deck.slide_width = Inches(13.333)
deck.slide_height = Inches(7.5)
BLANK = deck.slide_layouts[6]

DECK_NAME = "F1 Race Strategy Simulator"


def text(slide, value, x, y, w, h, size=16, bold=False, color=TEXT, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
    """Add a text box. `value` is a string or a list of paragraphs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor

    paragraphs = value if isinstance(value, list) else [value]
    for index, content in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        if spacing is not None and index > 0:
            paragraph.space_before = Pt(spacing)
        run = paragraph.add_run()
        run.text = content
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def label(slide, value, x, y, w=6.0, color=RED):
    """Small capital label above an equation or a block."""
    return text(slide, value.upper(), x, y, w, 0.26, size=11, bold=True, color=color,
                font=HEAD)


def panel(slide, x, y, w, h, fill=PANEL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = min(0.08, 0.14 / max(h, 0.1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def block(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def number_dot(slide, number, x, y, size=0.46):
    """The deck's motif: a red circle carrying a step number."""
    dot = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    dot.fill.solid()
    dot.fill.fore_color.rgb = RED
    dot.line.fill.background()
    dot.shadow.inherit = False
    frame = dot.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = str(number)
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.name = HEAD
    run.font.color.rgb = PAPER


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def content_slide(kicker, title, notes):
    """Light slide with the shared frame: kicker, title, footer, number."""
    slide = deck.slides.add_slide(BLANK)
    background(slide, PAPER)
    text(slide, kicker.upper(), 0.6, 0.42, 9, 0.3, size=12, bold=True, color=RED, font=HEAD)
    text(slide, title, 0.6, 0.74, 12.1, 0.75, size=30, bold=True, color=INK, font=HEAD,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, DECK_NAME, 0.6, 6.98, 6, 0.26, size=10, color=MUTED)
    text(slide, str(len(deck.slides)), 11.73, 6.98, 1, 0.26, size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def picture(slide, path, x, y, w):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))


# -------------------------------------------------------------- equations ---
EQ_DPI = 300
EQ_POINTS = 24.0
plt.rcParams["mathtext.fontset"] = "cm"


def render_equation(tex, color):
    """Typeset one equation to a transparent PNG; return path and size in inches."""
    rgb = "#%02X%02X%02X" % (color[0], color[1], color[2])
    name = hashlib.md5(f"{tex}|{rgb}".encode()).hexdigest()[:16] + ".png"
    path = equations / name
    if not path.exists():
        figure = plt.figure(figsize=(0.1, 0.1))
        figure.text(0, 0, f"${tex}$", fontsize=EQ_POINTS, color=rgb)
        figure.savefig(path, dpi=EQ_DPI, transparent=True, bbox_inches="tight",
                       pad_inches=0.03)
        plt.close(figure)
    with Image.open(path) as image:
        width, height = image.size
    return path, width / EQ_DPI, height / EQ_DPI


def eq(slide, tex, x, y, size=20, color=INK, max_w=None, box_w=None, box_h=None):
    """
    Place an equation with its left top at (x, y).

    `size` is the font size in points. With box_w / box_h the equation is
    centred in that width / height. max_w shrinks it to fit.
    """
    path, width, height = render_equation(tex, color)
    scale = size / EQ_POINTS
    limit = max_w if max_w is not None else box_w
    if limit is not None and width * scale > limit:
        scale = limit / width
    width, height = width * scale, height * scale
    if box_w is not None:
        x += (box_w - width) / 2
    if box_h is not None:
        y += (box_h - height) / 2
    shape = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(width))
    shape.name = "Equation"
    shape._element.nvPicPr.cNvPr.set("descr", tex)
    return width, height


def table(slide, rows, x, y, w, row_h, widths, size=14, numeric_from=1):
    """Table with a dark header row. `widths` are relative column widths."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                   Inches(w), Inches(row_h * len(rows)))
    grid = shape.table
    for c, share in enumerate(widths):
        grid.columns[c].width = Inches(w * share / sum(widths))
    for r, row in enumerate(rows):
        grid.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = grid.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK if r == 0 else (PANEL if r % 2 else PAPER)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if c < numeric_from else PP_ALIGN.RIGHT
            run = paragraph.add_run()
            run.text = value
            run.font.size = Pt(size)
            run.font.name = BODY
            run.font.bold = r == 0
            run.font.color.rgb = PAPER if r == 0 else TEXT
    return grid


def share_bar(slide, x, y, w, h, shares, names):
    """One horizontal bar split into labelled parts."""
    fills = [INK, MID, GREY]
    inks = [PAPER, PAPER, INK]
    left = x
    for share, name, fill, ink in zip(shares, names, fills, inks):
        width = w * share
        block(slide, left, y, width - 0.03, h, fill)
        if width > 0.9:
            text(slide, f"{name}  {share:.0%}", left + 0.12, y, width - 0.2, h, size=13,
                 bold=True, color=ink, anchor=MSO_ANCHOR.MIDDLE)
        left += width


def stat(slide, x, y, w, h, value, caption, value_color=INK, value_size=28):
    panel(slide, x, y, w, h)
    text(slide, value, x + 0.22, y + 0.12, w - 0.44, 0.6, size=value_size, bold=True,
         color=value_color, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, caption, x + 0.22, y + 0.78, w - 0.44, h - 0.85, size=13, color=MUTED)


# ---------------------------------------------------------------- numbers ---
sd = num["sd"]
var = num["variance"]
race = num["race_time"]
ex = num["example"]
pit = num["pit_lap"]
complexity = {row["leaves"]: row for row in num["complexity"]}
best_val = validation.set_index("model")

sigma_a, sigma_b, sigma_e = sd["race_sd_pct"], sd["stint_sd_pct"], sd["lap_sd_pct"]
mae_model = backtest["Race-time MAE on green races (s)"]
mae_naive = backtest["Race-time MAE with no model, constant baseline pace (s)"]
mae_pct = backtest["Race-time MAE on green races (% of race time)"]
pit_mae = backtest["First pit lap MAE vs team (laps)"]

# ================================================================ slides ===

# 1 ---------------------------------------------------------------- title ---
slide = deck.slides.add_slide(BLANK)
background(slide, INK)
text(slide, "PREDICTIVE MODELLING · COURSE PROJECT", 0.7, 1.3, 6, 0.3, size=12,
     bold=True, color=RED, font=HEAD)
text(slide, "Can a model call the pit stop?", 0.7, 1.75, 5.6, 2.1, size=46, bold=True,
     color=PAPER, font=HEAD)
text(slide, "A statistical model of lap time, the optimisation built on it, and "
     "whether its error allows a decision", 0.7, 4.0, 5.3, 1.1, size=18, color=SOFT_WHITE)
eq(slide, r"\hat{s} \;=\; \underset{s \,\in\, \mathcal{S}}{\mathrm{arg\,min}}\;\; \hat{R}(s)",
   0.7, 5.3, size=26, color=PAPER)
picture(slide, reports / "05_app" / "race_view.png", 6.65, 1.9, 6.1)
slide.notes_slide.notes_text_frame.text = (
    "The project has two mathematical parts. First a regression problem: predict how a "
    "driver's lap time changes through a race. Second an optimisation: choose the pit "
    "strategy that minimises predicted race time. The question is whether the error of "
    "the first part is small enough for the second part to mean anything."
)

# 2 -------------------------------------------------------------- problem ---
slide = content_slide(
    "Problem",
    "Pit strategy as an optimisation over a learned function",
    "A strategy s is a list of pit laps and tyre compounds. Its predicted race time is "
    "the sum of predicted lap times plus a fixed pit loss per stop. Each lap time is the "
    "driver's baseline pace B scaled by a learned function f-hat of the inputs on that "
    "lap, and those inputs depend on the strategy through tyre compound and tyre age. "
    "We want the strategy in the feasible set that minimises this. That raises two "
    "separate questions, at the bottom: how accurate is f-hat, and is the minimiser "
    "stable under its error. The answers differ.",
)
panel(slide, 0.6, 1.85, 6.9, 3.45, fill=INK)
label(slide, "Objective", 0.95, 2.08)
eq(slide, r"\hat{s} \;=\; \underset{s \,\in\, \mathcal{S}}{\mathrm{arg\,min}}\;\; \hat{R}(s)",
   0.6, 2.45, size=27, color=PAPER, box_w=6.9)
label(slide, "Predicted race time", 0.95, 3.55)
eq(slide, r"\hat{R}(s) \;=\; \sum_{\ell=\ell_0}^{N} B\left(1+\dfrac{\hat{f}\,(x_\ell(s))}{100}\right)"
   r"\;+\; k(s)\,P", 0.6, 3.9, size=23, color=PAPER, box_w=6.9, max_w=6.3)

panel(slide, 7.7, 1.85, 5.03, 3.45)
label(slide, "Notation", 8.0, 2.08, color=INK)
notation = [
    (r"s", "strategy: pit laps and tyre compounds"),
    (r"B", "baseline pace of the driver, seconds"),
    (r"\hat{f}", "learned model: % change in pace"),
    (r"x_\ell(s)", "inputs on lap ℓ under strategy s"),
    (r"k(s),\; P", "number of stops, time lost per stop"),
    (r"\mathcal{S}", "feasible strategies (tyre rules)"),
]
for index, (symbol, meaning) in enumerate(notation):
    ny = 2.5 + index * 0.45
    eq(slide, symbol, 8.0, ny, size=16, box_h=0.36)
    text(slide, meaning, 9.15, ny, 3.45, 0.36, size=14, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)

questions = [
    ("Statistics: how accurate is the estimate?",
     f"Race time within {mae_pct:.2f}% on an unseen season."),
    ("Decision: is the minimiser stable?",
     f"No. The recommended pit lap is {pit_mae:.0f} laps from the team's on average."),
]
for index, (question, answer) in enumerate(questions):
    qx = 0.6 + index * 6.16
    panel(slide, qx, 5.5, 5.97, 1.2)
    number_dot(slide, index + 1, qx + 0.25, 5.87)
    text(slide, question, qx + 0.95, 5.62, 4.8, 0.4, size=15, bold=True, color=INK,
         font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, answer, qx + 0.95, 6.05, 4.8, 0.55, size=14, color=MUTED)

# 3 -------------------------------------------------------- target and data ---
slide = content_slide(
    "Target and data",
    "The response is a ratio, so the circuit and car scale cancel",
    "Absolute lap time is dominated by the circuit and the car. Dividing by the "
    "driver's own baseline removes that multiplicative scale and leaves the quantity a "
    "strategy depends on: the percentage change in pace. The baseline is the median of "
    "up to eight clean laps between laps 2 and 12, so it is known early and no future "
    "information enters. The model is additive on this scale: y equals f of x plus an "
    "error. All nine inputs are known in advance for any hypothetical strategy, which "
    "is what allows the model to be used as a simulator. Data are real race laps from "
    "FastF1, split by season.",
)
panel(slide, 0.6, 1.85, 6.0, 4.85, fill=INK)
label(slide, "Response", 0.95, 2.05)
eq(slide, r"y_\ell \;=\; 100\left(\dfrac{T_\ell}{B}-1\right)", 0.95, 2.35, size=24, color=PAPER)
label(slide, "Baseline", 0.95, 3.35)
eq(slide, r"B \;=\; \mathrm{median}\,\{\,T_\ell \;:\; \ell \in \mathcal{B}\,\},\qquad"
   r" \mathcal{B} \subset \{2,\dots,12\},\;\; |\mathcal{B}| \leq 8",
   0.95, 3.68, size=17, color=PAPER, max_w=5.3)
label(slide, "Model and its inverse", 0.95, 4.35)
eq(slide, r"y_\ell = f(x_\ell) + e_\ell \qquad\quad \hat{T}_\ell = B\left(1+\dfrac{\hat{f}(x_\ell)}{100}\right)",
   0.95, 4.68, size=18, color=PAPER, max_w=5.3)
text(slide, "B is fixed a few laps into the race, so predicting later laps from it "
     "uses no future information.", 0.95, 5.65, 5.3, 0.8, size=14, color=SOFT_WHITE)

panel(slide, 6.9, 1.85, 5.83, 2.55)
eq(slide, r"x \in \mathbb{R}^6 \times \{\mathrm{categories}\}^3", 7.2, 2.03, size=17)
inputs_left = ["Tyre age a (laps)", "Fuel burned since baseline / N", "Race progress ℓ / N",
               "Race length N", "Baseline tyre age"]
inputs_right = ["Baseline track temperature", "Compound c: soft, medium, hard",
                "Baseline compound", "Circuit"]
text(slide, inputs_left, 7.2, 2.62, 2.7, 1.7, size=14, color=TEXT, spacing=3)
text(slide, inputs_right, 9.85, 2.62, 2.8, 1.7, size=14, color=TEXT, spacing=3)

table(slide, [("Split", "Seasons", "Dry races", "Laps n"),
              ("Train", "2022, 2023", "35", "24,039"),
              ("Validation", "2024", "19", "15,332"),
              ("Test", "2025", "21", "15,844")],
      6.9, 4.6, 5.83, 0.4, [1.2, 1.3, 1.0, 1.0], size=14, numeric_from=2)
text(slide, "Split by season, not at random: errors within one race are correlated "
     f"(ρ ≈ {var['share_race']:.2f}, slide 7), so a random split would leak.",
     6.9, 6.3, 5.83, 0.5, size=13, color=MUTED)

# 4 ---------------------------------------------------- reference and linear ---
slide = content_slide(
    "Reference and linear models",
    "Constant pace is biased; a linear model fixes most of it",
    "Three steps. The naive reference predicts zero change. Its mean squared error "
    "decomposes exactly into squared bias plus variance of the response: cars get faster "
    "as fuel burns, the mean change is minus 1.18 percent, and that bias is why its "
    "R squared is negative. Ordinary least squares with six standardised numeric inputs "
    "and indicator variables for compound gives 0.930 on validation with almost no gap "
    "to training error, so it is a low-variance model. Adding a quadratic in tyre age "
    "and compound interactions, with a ridge penalty, does not help.",
)
cards = [
    ("No change (predict zero)",
     r"\mathrm{MSE} = \bar{y}^{\,2} + s_y^2:\qquad %.3f^2 = %.3f^2 + %.3f^2"
     % (num["val_rmse_zero"], abs(num["val_mean_y"]), num["val_sd_y"]),
     f"Fuel burn makes cars faster: mean change {abs(num['val_mean_y']):.2f}% down. That bias "
     f"gives R² = −{abs(best_val.loc[NAIVE, 'r2']):.2f}."),
    ("Linear regression (least squares)",
     r"\hat{\beta} = \underset{\beta}{\mathrm{arg\,min}}\; \|y - X\beta\|^2"
     r" = (X^{\mathrm{T}}X)^{-1}X^{\mathrm{T}}y",
     f"{num['linear_coefficients']} coefficients. RMSE: training {num['train_rmse_linear']:.3f}%, "
     f"validation {num['val_rmse_linear']:.3f}%. Almost no gap: low variance."),
    ("Wear curve per compound (ridge)",
     r"\min_{\beta}\; \|y - X\beta\|^2 + \alpha\,\|\beta\|^2,\qquad \alpha = 1",
     "Adds a² and a × compound terms. Validation "
     f"{best_val.loc['Linear, wear per compound', 'rmse_pct']:.3f}%: the extra terms "
     "do not help."),
]
for index, (name, formula, detail) in enumerate(cards):
    cy = 1.85 + index * 1.63
    panel(slide, 0.6, cy, 6.55, 1.5)
    label(slide, name, 0.88, cy + 0.14, color=INK)
    eq(slide, formula, 0.88, cy + 0.42, size=16, box_h=0.5, max_w=6.0)
    text(slide, detail, 0.88, cy + 0.96, 6.05, 0.5, size=13, color=MUTED)

short = {"Gradient boosting (monotonic tyre wear)": "Boosting, monotone wear",
         "Linear, wear per compound": "Linear, wear per compound",
         NAIVE: "No change (ŷ = 0)", "Mean change": "Mean change (ŷ = ȳ)"}
rows = [("Model", "RMSE %", "RMSE s", "R²")]
for _, row in validation.sort_values("rmse_pct").iterrows():
    rows.append((short.get(row["model"], row["model"]), f"{row['rmse_pct']:.3f}",
                 f"{row['rmse_s']:.3f}", f"{row['r2']:.3f}".replace("-", "−")))
text(slide, "Validation season 2024", 7.45, 1.85, 5.28, 0.35, size=15, bold=True,
     color=INK, font=HEAD)
table(slide, rows, 7.45, 2.3, 5.28, 0.48, [2.9, 1.0, 1.0, 1.0], size=14)
text(slide, "RMSE in % of baseline pace and in seconds per lap. Models were fitted on "
     "2022–2023 and compared here; 2025 was not touched.", 7.45, 5.85, 5.28, 0.75,
     size=13, color=MUTED)

# 5 -------------------------------------------------------------- boosting ---
slide = content_slide(
    "Gradient boosting",
    "Boosting: gradient descent in function space",
    "Gradient boosting builds the predictor as a sum of small regression trees. Start "
    "from the mean. At each step compute the negative gradient of the loss at the "
    "current fit; for squared loss that is just the residual. Fit a tree with J leaves "
    "to those residuals and add a fraction nu of it. Each leaf value is a shrunken mean "
    "of the residuals in the leaf, with an L2 penalty lambda in the denominator, and "
    "splits are chosen to maximise the gain shown. The chart is the fit stage by stage: "
    "validation error falls below linear regression after roughly a hundred trees and "
    "then flattens. Ten percent of training laps were held out internally for early "
    "stopping; it never triggered, so all 300 trees are used.",
)
panel(slide, 0.6, 1.85, 6.25, 4.85, fill=INK)
steps = [
    ("Additive model", r"F_0(x)=\bar{y},\qquad F_m(x) = F_{m-1}(x) + \nu\, h_m(x)", 19),
    ("Pseudo-residuals for squared loss",
     r"r_{im} \;=\; -\,\dfrac{\partial L}{\partial F}\,(x_i) \;=\; y_i - F_{m-1}(x_i),"
     r"\qquad L=\frac{1}{2}(y-F)^2", 17),
    ("Leaf value",
     r"w_j \;=\; \dfrac{1}{n_j+\lambda}\sum_{i\in R_j} r_{im}", 18),
    ("Split gain",
     r"\dfrac{1}{2}\left[\dfrac{G_L^2}{n_L+\lambda}+\dfrac{G_R^2}{n_R+\lambda}"
     r"-\dfrac{G^2}{n+\lambda}\right],\qquad G=\sum_i r_{im}", 17),
]
for index, (name, formula, size) in enumerate(steps):
    sy = (2.02, 3.02, 4.1, 5.32)[index]
    label(slide, name, 0.95, sy)
    eq(slide, formula, 0.95, sy + 0.3, size=size, color=PAPER, box_h=0.78, max_w=5.6)

picture(slide, assets / "fig_boosting.png", 7.1, 1.85, 5.63)
chips = [r"M = 300", r"\nu = 0.03", r"J = 4", r"\lambda = 5", r"n_j \geq 300"]
for index, chip in enumerate(chips):
    cx = 7.1 + index * 1.14
    panel(slide, cx, 5.4, 1.05, 0.62)
    eq(slide, chip, cx, 5.4, size=14, box_w=1.05, box_h=0.62, max_w=0.9)
text(slide, "Trees, learning rate, leaves per tree, L2 penalty, minimum leaf size. "
     "The chart is the error after each tree.", 7.1, 6.15, 5.63, 0.55, size=13, color=MUTED)

# 6 ------------------------------------------------------------ complexity ---
slide = content_slide(
    "Model selection",
    "Tree size is chosen by the bias–variance trade-off",
    "Expected squared error splits into squared bias, variance and irreducible noise. "
    "Leaves per tree control the trade. As J grows, training error keeps falling, but "
    "validation error is lowest at J equals 4 and then rises: with 31 leaves the "
    "boosted model is worse than linear regression. The reason is on the next slide: "
    "the data have far fewer independent observations than laps. The selected model "
    "was refitted on 2022 to 2024 and applied once to 2025: 0.835 seconds per lap "
    "against 1.558 for constant pace, R squared 0.44.",
)
panel(slide, 0.6, 1.85, 6.4, 0.95)
eq(slide, r"\mathbb{E}\left[(y-\hat{f}(x))^2\right] \;=\; \mathrm{Bias}[\hat{f}(x)]^2"
   r" \;+\; \mathrm{Var}[\hat{f}(x)] \;+\; \sigma^2", 0.6, 1.85, size=18, box_w=6.4,
   box_h=0.95, max_w=5.9)
picture(slide, assets / "fig_complexity.png", 0.6, 2.95, 6.3)

rows = [("Leaves J", "Training", "Validation", "Gap")]
for leaves in (2, 4, 8, 31):
    row = complexity[leaves]
    rows.append((str(leaves), f"{row['train']:.3f}", f"{row['validation']:.3f}",
                 f"{row['validation'] - row['train']:.3f}"))
rows.append(("Linear", f"{num['train_rmse_linear']:.3f}", f"{num['val_rmse_linear']:.3f}",
             f"{num['val_rmse_linear'] - num['train_rmse_linear']:.3f}"))
text(slide, "RMSE, % of baseline pace", 7.35, 1.85, 5.38, 0.35, size=15, bold=True,
     color=INK, font=HEAD)
table(slide, rows, 7.35, 2.28, 5.38, 0.4, [1.2, 1.0, 1.0, 1.0], size=14)
text(slide, "Training error falls with J; validation error is lowest at J = 4. "
     "The growing gap is variance.", 7.35, 4.78, 5.38, 0.6, size=14, color=TEXT)

text(slide, "Test season 2025, used once", 7.35, 5.42, 5.38, 0.3, size=13, bold=True,
     color=INK, font=HEAD)
tests = [(f"{test.loc[MODEL, 'rmse_s']:.3f} s", "RMSE per lap", RED),
         (f"{test.loc[NAIVE, 'rmse_s']:.3f} s", "constant pace", INK),
         (f"{test.loc[MODEL, 'r2']:.2f}", "R²", INK)]
for index, (value, caption, color) in enumerate(tests):
    tx = 7.35 + index * 1.83
    panel(slide, tx, 5.78, 1.72, 0.95)
    text(slide, value, tx + 0.15, 5.83, 1.45, 0.48, size=20, bold=True, color=color,
         font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, caption, tx + 0.15, 6.32, 1.45, 0.3, size=12, color=MUTED)

# 7 --------------------------------------------------------- decomposition ---
slide = content_slide(
    "Error structure",
    "The errors are not independent: a three-level decomposition",
    "This is the central modelling step. The validation residual of driver-race d, "
    "stint j, lap l is split into a term shared by the whole race, a term shared by the "
    "stint, and a lap term. About a third of the error variance is the race-level term, "
    "which is mostly a misjudged baseline. The intraclass correlation is therefore 0.33. "
    "With 43 laps per driver-race the design effect is about 15, so 15 thousand "
    "validation laps carry the information of roughly one thousand independent ones. "
    "That is why the split must be by season and why small trees win. The three terms "
    "also behave differently in a strategy comparison, which the next slides use. The "
    "components are standard deviations of group means; they sum to slightly more than "
    "the pooled mean squared error.",
)
panel(slide, 0.6, 1.85, 5.6, 2.3, fill=INK)
label(slide, "Validation residual, by level", 0.95, 2.05)
eq(slide, r"e_{dj\ell} \;=\; a_d \;+\; b_{dj} \;+\; \varepsilon_{dj\ell}", 0.6, 2.4,
   size=27, color=PAPER, box_w=5.6)
eq(slide, r"a_d \sim (0,\sigma_a^2) \qquad b_{dj} \sim (0,\sigma_b^2) \qquad"
   r" \varepsilon_{dj\ell} \sim (0,\sigma_\varepsilon^2)", 0.6, 3.3, size=16,
   color=SOFT_WHITE, box_w=5.6, max_w=5.0)
text(slide, "d: driver-race,  j: stint,  ℓ: lap.  Mutually independent, mean zero.", 0.6, 3.75, 5.6, 0.3, size=12,
     color=SOFT_WHITE, align=PP_ALIGN.CENTER)

levels = [
    (r"\sigma_a = %.3f" % sigma_a, "Driver over a race",
     "A misjudged baseline. The same for every strategy."),
    (r"\sigma_b = %.3f" % sigma_b, "Stint within a race",
     "The part that separates strategies."),
    (r"\sigma_\varepsilon = %.3f" % sigma_e, "Lap within a stint",
     "Averages out over a stint."),
]
for index, (formula, name, detail) in enumerate(levels):
    ly = 1.85 + index * 0.79
    panel(slide, 6.45, ly, 6.28, 0.72)
    eq(slide, formula, 6.7, ly, size=16, box_h=0.72)
    text(slide, name, 8.75, ly + 0.08, 3.8, 0.28, size=14, bold=True, color=INK, font=HEAD)
    text(slide, detail, 8.75, ly + 0.37, 3.8, 0.3, size=13, color=MUTED)

text(slide, "Share of error variance (in % of pace, squared)", 0.6, 4.38, 8, 0.3,
     size=13, bold=True, color=INK, font=HEAD)
share_bar(slide, 0.6, 4.72, 12.13, 0.5,
          [var["share_race"], var["share_stint"], var["share_lap"]],
          ["Race", "Stint", "Lap"])

panel(slide, 0.6, 5.42, 5.97, 1.3)
label(slide, "Intraclass correlation", 0.88, 5.52, color=INK)
eq(slide, r"\rho \;=\; \dfrac{\sigma_a^2}{\sigma_a^2+\sigma_b^2+\sigma_\varepsilon^2} \;=\; %.2f"
   % var["share_race"], 0.88, 5.8, size=17, box_h=0.85)
panel(slide, 6.76, 5.42, 5.97, 1.3)
label(slide, "Effective sample size", 7.04, 5.52, color=INK)
eq(slide, r"n_{\mathrm{eff}} \;=\; \dfrac{n}{1+(\bar{m}-1)\,\rho} \;=\; \dfrac{15{,}332}{%.1f}"
   r" \;\approx\; %s" % (var["design_effect"],
                         f"{round(var['effective_laps'], -1):,.0f}".replace(",", "{,}")),
   7.04, 5.8, size=17, box_h=0.85, max_w=5.4)

# 8 ----------------------------------------------------------- propagation ---
N_demo, stints_demo = 50, (25, 25)
terms = [N_demo ** 2 * var["race"], sum(n * n for n in stints_demo) * var["stint"],
         N_demo * var["lap"]]
term_shares = [value / sum(terms) for value in terms]

slide = content_slide(
    "From laps to races",
    "Propagating the error predicts the race-time error",
    "Race time is baseline times the sum of one plus y over a hundred, so the race-time "
    "error is B over a hundred times the sum of lap errors. Under the three-level model "
    "its variance has three terms. The race-level term is perfectly correlated across "
    "laps and grows with N squared; lap noise grows only with N. So lap noise, 46 "
    "percent of lap variance, is 2 percent of race-time variance, and the race-level "
    "term is three quarters. Now the check: the variances were estimated on 2024. "
    "Applied to the 60 green-flag driver-races of 2025, the formula predicts a "
    "standard deviation of 28.0 seconds. The observed root mean squared error is 27.9. "
    "The standardised errors have standard deviation 1.03, where 1 is what the model "
    "implies. The histogram is rough rather than bell-shaped: sixty driver-races from "
    "a limited number of events, and drivers in the same race share conditions. So the "
    "size of the error is predicted out of sample; its exact shape is not established.",
)
panel(slide, 0.6, 1.85, 6.6, 2.7, fill=INK)
label(slide, "Race-time error for a given strategy", 0.95, 2.05)
eq(slide, r"R-\hat{R} \;=\; \dfrac{B}{100}\sum_{\ell=\ell_0}^{N} e_\ell", 0.95, 2.35,
   size=19, color=PAPER)
label(slide, "Its variance", 0.95, 3.3)
eq(slide, r"\sigma_R^2 \;=\; \left(\dfrac{B}{100}\right)^{2}\left[\,N^2\sigma_a^2 \;+\;"
   r" \sigma_b^2\sum_j n_j^2 \;+\; N\sigma_\varepsilon^2\,\right]", 0.95, 3.62, size=19,
   color=PAPER, max_w=5.9)

text(slide, "Share of race-time variance (N = 50 laps, two stints of 25)", 0.6, 4.75,
     6.6, 0.3, size=13, bold=True, color=INK, font=HEAD)
share_bar(slide, 0.6, 5.08, 6.6, 0.5, term_shares, ["Race", "Stint", "Lap"])
text(slide, f"Lap noise is {var['share_lap']:.0%} of the lap-level variance but "
     f"{term_shares[2]:.0%} of race-time variance: it averages out. The race-level "
     "term adds up coherently and dominates.", 0.6, 5.75, 6.6, 0.95, size=14, color=TEXT)

picture(slide, assets / "fig_race_time_z.png", 7.7, 1.8, 4.75)
text(slide, f"{race['n']} green-flag driver-races, 2025. Curve: standard normal. The "
     "spread matches; the shape is rough.", 7.5, 4.72, 5.23, 0.5, size=12, color=MUTED)
checks = [(f"{race['implied_sd_s']:.1f} s", "predicted SD, from 2024 variances", RED),
          (f"{race['observed_rmse_s']:.1f} s", "observed RMSE, 2025", INK),
          (f"{race['z_sd']:.2f}", "SD of z (1 if the model holds)", INK)]
for index, (value, caption, color) in enumerate(checks):
    cx = 7.5 + index * 1.78
    panel(slide, cx, 5.35, 1.67, 1.35)
    text(slide, value, cx + 0.14, 5.42, 1.4, 0.5, size=20, bold=True, color=color,
         font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, caption, cx + 0.14, 5.95, 1.42, 0.7, size=12, color=MUTED)

# 9 ---------------------------------------------------------------- search ---
one_stop_laps = ex["one_stop_candidates"] // 3
two_stop_pairs = ex["two_stop_candidates"] // 9
evaluations = ex["valid_strategies"] * ex["n_laps"]

slide = content_slide(
    "Optimisation",
    "A small constrained search, solved exactly by enumeration",
    "The feasible set has three constraints: every stint at least five laps, at least "
    "two different compounds as the rules require, and no stint beyond the 98th "
    "percentile of tyre age seen on that compound at that circuit, because trees cannot "
    "extrapolate. The set is small enough to enumerate. For the example, the 2025 Saudi "
    "Arabian Grand Prix winner, there are 90 one-stop and 702 two-stop candidates, 497 "
    "feasible, evaluated in one batched model call. So the minimiser is exact for the "
    "model: any error in the recommendation comes from R-hat, not from the search. The "
    "chart shows the problem. The objective is flat: 15 strategies are within two "
    "seconds of the best, while one standard deviation of the predicted time is 12 "
    "seconds.",
)
panel(slide, 0.6, 1.85, 6.1, 1.75, fill=INK)
label(slide, "Feasible set", 0.95, 2.05)
eq(slide, r"\mathcal{S}=\{\,s:\;\; n_j\geq 5,\;\;\; |\{c_j\}|\geq 2,\;\;\;"
   r" a_j\leq A_{0.98}(c_j)\;\;\; \forall j\,\}", 0.6, 2.4, size=18, color=PAPER,
   box_w=6.1, box_h=0.6, max_w=5.5)
text(slide, "n: stint length.  c: compound.  a: tyre age at the end of the stint.",
     0.95, 3.12, 5.5, 0.3, size=12, color=SOFT_WHITE)

text(slide, f"Example: {ex['event']}, winner", 0.6, 3.8, 6.1, 0.3, size=13, bold=True,
     color=INK, font=HEAD)
counts = [
    (r"3 \times %d = %d" % (one_stop_laps, ex["one_stop_candidates"]),
     "one-stop candidates, 3 compounds"),
    (r"9 \times %d = %d" % (two_stop_pairs, ex["two_stop_candidates"]),
     "two-stop candidates, pit-lap step 2"),
    (r"|\mathcal{S}| = %d" % ex["valid_strategies"],
     f"feasible: {ex['valid_one_stop']} one-stop, {ex['valid_two_stop']} two-stop"),
    (r"%d \times %d = %s" % (ex["valid_strategies"], ex["n_laps"],
                             f"{evaluations:,}".replace(",", "{,}")),
     "lap predictions, one model call"),
]
for index, (formula, detail) in enumerate(counts):
    cy = 4.17 + index * 0.64
    panel(slide, 0.6, cy, 6.1, 0.57)
    eq(slide, formula, 0.82, cy, size=14, box_h=0.57, max_w=2.2)
    text(slide, detail, 2.95, cy, 3.65, 0.57, size=13, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

picture(slide, assets / "fig_objective.png", 6.95, 1.8, 5.78)
text(slide, f"The objective is flat near its minimum: {ex['within_2s_of_best']} "
     f"strategies are within 2 s of the best. The grey band is one standard deviation "
     f"of a predicted race time ({ex['noise_sd_best_s']:.0f} s).", 6.95, 5.5, 5.78, 0.9,
     size=14, color=TEXT)
text(slide, "Enumeration is exact, so the only error in the recommendation is the "
     "error in the predicted race time.", 6.95, 6.3, 5.78, 0.5, size=13, color=MUTED)

# 10 ------------------------------------------------------------ stability ---
ratio = ex["gap_team_minus_best_s"] / ex["sd_gap_s"]

# The same comparison if the two strategies share their error on the laps
# where the car is in the same tyre state (the laps before the earlier stop).
early, late = sorted([ex["team_pit_lap"], ex["recommended_pit_lap"]])
pieces = [late - early, ex["total_laps"] - late, ex["total_laps"] - early]
sd_gap_shared = ex["baseline_s"] / 100 * math.sqrt(
    var["stint"] * sum(n * n for n in pieces) + var["lap"] * sum(pieces))
p_mc = ex["prob_best_faster_monte_carlo"]
se_mc = math.sqrt(p_mc * (1 - p_mc) / 4000)

slide = content_slide(
    "Decision under uncertainty",
    "The minimiser is not stable: the signal is a fifth of the noise",
    "When two strategies for the same driver are compared, the race-level error a "
    "cancels, because it shifts both equally. What remains is stint and lap error. For "
    "the example the recommended strategy is predicted 3.2 seconds faster than the "
    "team's, but the standard deviation of that difference is 18 seconds, treating the "
    "two races as independent draws; if the laps the two strategies have in common "
    "share their error it is 16 seconds, which changes nothing. The probability that the recommendation is really "
    "faster is Phi of 0.18, which is 0.57. A Monte Carlo of four thousand simulated "
    "races gives 0.578 with standard error 0.008, so simulation and formula agree. "
    "Barely better than a coin. Across all 182 driver-races the same thing shows up as "
    "disagreement with the teams: same number of stops in 65 percent, but the first pit "
    "lap is 8 laps away on average.",
)
panel(slide, 0.6, 1.85, 6.6, 3.0, fill=INK)
label(slide, "Noise separating two strategies (race term cancels)", 0.95, 2.05)
eq(slide, r"\sigma^2(s) \;=\; \left(\dfrac{B}{100}\right)^{2}\left[\,\sigma_b^2\sum_j n_j^2"
   r" \;+\; N\sigma_\varepsilon^2\,\right]", 0.95, 2.35, size=19, color=PAPER, max_w=5.9)
label(slide, "Probability the recommendation is really faster", 0.95, 3.45)
eq(slide, r"P \;=\; \Phi\left(\dfrac{\Delta}{\sigma_\Delta}\right),\qquad"
   r" \Delta=\hat{R}(s')-\hat{R}(\hat{s}),\quad"
   r" \sigma_\Delta^2=\sigma^2(\hat{s})+\sigma^2(s')", 0.95, 3.78, size=17, color=PAPER,
   max_w=5.9)

text(slide, f"Example: {ex['event']}, winner", 7.45, 1.85,
     5.28, 0.3, size=13, bold=True, color=INK, font=HEAD)
worked = [
    (r"\Delta = %.1f\ \mathrm{s}" % ex["gap_team_minus_best_s"], "predicted gain"),
    (r"\sigma_\Delta = %.1f\ \mathrm{s}" % ex["sd_gap_s"], "its standard deviation"),
    (r"P = \Phi(%.2f) = %.2f" % (ratio, ex["prob_best_faster_normal"]), "from the formula"),
    (r"\hat{P} = %.3f \pm %.3f" % (p_mc, se_mc), "Monte Carlo, M = 4000 races"),
]
for index, (formula, detail) in enumerate(worked):
    wy = 2.22 + index * 0.66
    panel(slide, 7.45, wy, 5.28, 0.59)
    eq(slide, formula, 7.67, wy, size=15, box_h=0.59, max_w=2.5)
    text(slide, detail, 10.3, wy, 2.35, 0.59, size=13, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

text(slide, "All 182 driver-races of 2025: recommendation against the team's choice",
     0.6, 5.05, 12.13, 0.3, size=13, bold=True, color=INK, font=HEAD)
outcomes = [
    (f"{backtest['Same number of stops as the team (%)']:.0f}%", "same number of stops"),
    (f"{pit_mae:.1f} laps", "mean distance of first pit lap"),
    (f"{backtest['First pit lap within 3 laps of team (%)']:.0f}%", "within 3 laps"),
    (f"{backtest['Median predicted gain of recommended strategy (s)']:.1f} s",
     "median predicted gain Δ"),
]
for index, (value, caption) in enumerate(outcomes):
    ox = 0.6 + index * 3.07
    panel(slide, ox, 5.4, 2.92, 1.0)
    text(slide, value, ox + 0.2, 5.45, 2.5, 0.5, size=22, bold=True,
         color=RED if index == 1 else INK, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, caption, ox + 0.2, 5.97, 2.5, 0.35, size=13, color=MUTED)
text(slide, "The standard deviation of Δ treats the two races as independent draws. "
     "Sharing the error on the laps both strategies have in common gives "
     f"{sd_gap_shared:.0f} s instead of {ex['sd_gap_s']:.0f} s.", 0.6, 6.5, 12.13, 0.3,
     size=12, color=MUTED)

# 11 ---------------------------------------------------------------- demo ---
slide = content_slide(
    "Demo",
    "The same mathematics as an animated race",
    "This is the app. The same driver runs three strategies at once on the real circuit "
    "outline: the minimiser s-hat, the team's real strategy, and one the user sets. "
    "Each lap is the model's prediction. With random variation switched on, each run "
    "draws one stint offset b per stint and one lap noise term per lap from the "
    "estimated distributions, so it is one Monte Carlo draw played out, and the "
    "finishing order changes between runs. Switch to the live app here.",
)
picture(slide, reports / "05_app" / "race_view.png", 0.6, 1.9, 8.0)
points = [
    (r"\hat{s},\; s',\; s^{\mathrm{you}}", "Three strategies race",
     "The minimiser, the team's, and your own"),
    (r"\hat{T}_\ell", "Every lap is a prediction",
     "Baseline pace scaled by the model"),
    (r"b_j,\; \varepsilon_\ell", "Random races",
     "One Monte Carlo draw of stint and lap error"),
    (r"\hat{R}(s)", "Live running order",
     "Cumulative predicted time, gap to the leader"),
]
for index, (symbol, name, detail) in enumerate(points):
    py = 1.9 + index * 1.13
    panel(slide, 8.9, py, 3.83, 1.0)
    eq(slide, symbol, 9.05, py, size=15, box_w=1.15, box_h=1.0, max_w=1.05)
    text(slide, name, 10.35, py + 0.14, 2.3, 0.3, size=14, bold=True, color=INK, font=HEAD)
    text(slide, detail, 10.35, py + 0.45, 2.3, 0.5, size=12, color=MUTED)
text(slide, "Circuit outlines are real. Start-line position and direction of travel "
     "are illustrative.", 0.6, 6.45, 8.0, 0.35, size=12, color=MUTED)

# 12 -------------------------------------------------------------- limits ---
slide = content_slide(
    "Limitations",
    "Where the assumptions fail",
    "Four failures, each stated as mathematics. First, selection: a lap at tyre age a "
    "is observed only if the stint lasted that long, and stints last when tyres hold "
    "up. So the observed mean at high age is below the true mean, f-hat is too "
    "optimistic about old tyres, and recommendations are later than the teams' in two "
    "thirds of cases, by 4.3 laps on average. Second, omitted variables: traffic and "
    "track position are not observed, which is why R squared is 0.44, and teams "
    "optimise for them. Third, safety cars change the pit loss and are not modelled; "
    "race time could only be checked on 60 of 182 driver-races. Fourth, resolution: "
    "the gains are far inside the error.",
)
limits = [
    ("Selection in tyre data",
     r"\mathbb{E}\,[\,y \mid a,\; \mathrm{stint} \geq a\,] \;\leq\; \mathbb{E}\,[\,y \mid a\,]",
     f"Old tyres are seen only when they lasted. Stops come out later than the team's "
     f"in {pit['share_later']:.0%} of cases, by {pit['mean_signed_laps']:.1f} laps on average."),
    ("Omitted variables",
     r"y \;=\; f(x) \;+\; g(u) \;+\; \varepsilon,\qquad u\ \mathrm{unobserved}",
     f"Traffic and track position are not in the data (R² = {test.loc[MODEL, 'r2']:.2f}). "
     "Teams stop early to gain position; the model sees clean air only."),
    ("A regime that is not modelled",
     r"P \;\to\; P_{\mathrm{SC}} \;<\; P \quad \mathrm{under\ a\ safety\ car}",
     f"Safety cars make a stop cheaper and decide many races. Race time could be "
     f"checked on {race['n']} of 182 driver-races only."),
    ("Resolution of the decision",
     r"|\Delta| \approx %.0f\ \mathrm{s} \;\ll\; \sigma_\Delta \approx %.0f\ \mathrm{s}"
     % (backtest["Median predicted gain of recommended strategy (s)"], ex["sd_gap_s"]),
     "Differences between good strategies are far inside the error. The team's choice "
     "is a reference, not a proven optimum."),
]
for index, (name, formula, detail) in enumerate(limits):
    lx = 0.6 + (index % 2) * 6.16
    ly = 1.85 + (index // 2) * 2.02
    panel(slide, lx, ly, 5.97, 1.88)
    number_dot(slide, index + 1, lx + 0.25, ly + 0.22, size=0.42)
    text(slide, name, lx + 0.85, ly + 0.22, 4.9, 0.42, size=15, bold=True, color=INK,
         font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    eq(slide, formula, lx + 0.85, ly + 0.72, size=15, box_h=0.42, max_w=4.9)
    text(slide, detail, lx + 0.85, ly + 1.22, 4.9, 0.6, size=13, color=MUTED)
panel(slide, 0.6, 5.92, 12.13, 0.8, fill=INK)
text(slide, "Next: a survival model for stint length, a safety-car scenario with its "
     "own pit loss, and a traffic term.", 0.95, 5.92, 11.4, 0.8, size=15, color=PAPER,
     anchor=MSO_ANCHOR.MIDDLE)

# 13 ------------------------------------------------------------- closing ---
slide = deck.slides.add_slide(BLANK)
background(slide, INK)
text(slide, "CONCLUSION", 0.9, 1.1, 6, 0.3, size=12, bold=True, color=RED, font=HEAD)
text(slide, "An accurate model is not yet a decision", 0.9, 1.5, 11.5, 1.0, size=40,
     bold=True, color=PAPER, font=HEAD)
closing = [
    ("Lap", r"\mathrm{RMSE} = %.3f\ \mathrm{s},\quad R^2 = %.2f"
     % (test.loc[MODEL, "rmse_s"], test.loc[MODEL, "r2"]),
     "Unseen season. Constant pace: "
     f"{test.loc[NAIVE, 'rmse_s']:.3f} s."),
    ("Race", r"\sigma_R:\ \ %.1f\ \mathrm{s\ predicted},\ \ %.1f\ \mathrm{s\ observed}"
     % (race["implied_sd_s"], race["observed_rmse_s"]),
     "The error model holds out of sample."),
    ("Decision", r"\Delta\,/\,\sigma_\Delta \approx %.1f \;\;\Rightarrow\;\; P \approx %.2f"
     % (ratio, ex["prob_best_faster_normal"]),
     "The minimiser is not identified."),
]
for index, (name, formula, detail) in enumerate(closing):
    cy = 2.85 + index * 0.98
    number_dot(slide, index + 1, 0.9, cy + 0.1)
    text(slide, name, 1.6, cy, 1.6, 0.66, size=18, bold=True, color=PAPER, font=HEAD,
         anchor=MSO_ANCHOR.MIDDLE)
    eq(slide, formula, 3.2, cy, size=19, color=PAPER, box_h=0.66, max_w=5.2)
    text(slide, detail, 8.75, cy, 3.9, 0.66, size=15, color=SOFT_WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
text(slide, "Validate a model at the level of the decision it will be used for.",
     0.9, 6.05, 11.5, 0.5, size=20, bold=True, color=PAPER, font=HEAD)
slide.notes_slide.notes_text_frame.text = (
    "Three levels, three answers. At lap level the model halves the error of constant "
    "pace. At race level the variance model predicts the size of the error almost "
    "exactly on a new season. At decision level the predicted gain is a fifth of its "
    "standard deviation, so the argmin is not identified by the data. Accuracy has to "
    "be judged at the level of the decision. Questions."
)

output.parent.mkdir(parents=True, exist_ok=True)
deck.save(output)
print("saved", output, "slides:", len(deck.slides))
